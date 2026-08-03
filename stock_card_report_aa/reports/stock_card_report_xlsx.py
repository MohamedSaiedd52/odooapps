from datetime import datetime

from odoo import fields, models

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GREY = Side(style="thin", color="D3D3D3")
BORDER = Border(left=GREY, right=GREY, top=GREY, bottom=GREY)
HEADER_FILL = PatternFill("solid", fgColor="CCFFFF")
TITLE_FONT = Font(bold=True, size=14)
HEADER_FONT = Font(bold=True)
CENTER = Alignment(horizontal="center")
LEFT = Alignment(horizontal="left")
RIGHT = Alignment(horizontal="right")
NUM_FORMAT = "#,##0.00"
DATE_FORMAT = "YYYY-MM-DD"

COLUMNS = ["Date", "Reference", "In", "Out", "Balance"]
COLUMN_WIDTHS = [22, 40, 14, 14, 14]


class ReportStockCardReportXlsx(models.AbstractModel):
    _name = "report.stock_card_report_aa.report_stock_card_report_aa_xlsx"
    _description = "Stock Card Report XLSX"
    _inherit = "report.report_xlsx_aa_aa.abstract"

    def generate_xlsx_report(self, workbook, data, objects):
        for product in objects.product_ids:
            ws = workbook.create_sheet(self._check_ws_name(product.name))
            self._write_product_sheet(ws, objects, product)

    def _setup_page(self, ws):
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        report_date = fields.Datetime.context_timestamp(
            self.env.user, datetime.now()
        ).strftime("%Y-%m-%d %H:%M")
        ws.oddFooter.left.text = report_date
        ws.oddFooter.left.size = 8
        ws.oddFooter.right.text = "&P / &N"
        ws.oddFooter.right.size = 8
        for col, width in enumerate(COLUMN_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width

    def _write_header_cell(self, ws, row, col, value):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = CENTER
        return cell

    def _write_data_cell(self, ws, row, col, value, align=RIGHT, num_format=None):
        cell = ws.cell(row=row, column=col, value=value)
        cell.border = BORDER
        cell.alignment = align
        if num_format:
            cell.number_format = num_format
        return cell

    def _write_product_sheet(self, ws, report, product):
        self._setup_page(ws)

        # Title
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
        title = ws.cell(row=1, column=1, value="Stock Card - {}".format(product.name))
        title.font = TITLE_FONT

        # Filters table
        self._write_header_cell(ws, 3, 1, "Date from")
        self._write_header_cell(ws, 3, 2, "Date to")
        self._write_header_cell(ws, 3, 3, "Location")
        self._write_data_cell(
            ws, 4, 1, report.date_from, align=CENTER, num_format=DATE_FORMAT
        )
        self._write_data_cell(
            ws, 4, 2, report.date_to, align=CENTER, num_format=DATE_FORMAT
        )
        self._write_data_cell(
            ws, 4, 3, report.location_id.display_name or "", align=CENTER
        )

        # Stock card table header
        header_row = 6
        for col, label in enumerate(COLUMNS, start=1):
            self._write_header_cell(ws, header_row, col, label)
        ws.freeze_panes = "A{}".format(header_row + 1)

        # Initial balance line
        row = header_row + 1
        balance = report._get_initial(
            report.results.filtered(
                lambda l: l.product_id == product and l.is_initial
            )
        )
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        self._write_data_cell(ws, row, 1, "Initial", align=CENTER)
        for col in range(2, 5):
            ws.cell(row=row, column=col).border = BORDER
        self._write_data_cell(ws, row, 5, balance, num_format=NUM_FORMAT)

        # Movement lines
        product_lines = report.results.filtered(
            lambda l: l.product_id == product and not l.is_initial
        )
        for line in product_lines:
            row += 1
            balance += line.product_in - line.product_out
            date_value = (
                fields.Datetime.context_timestamp(self.env.user, line.date).replace(
                    tzinfo=None
                )
                if line.date
                else ""
            )
            self._write_data_cell(
                ws, row, 1, date_value, align=LEFT, num_format=DATE_FORMAT
            )
            self._write_data_cell(ws, row, 2, line.display_name or "", align=LEFT)
            self._write_data_cell(
                ws, row, 3, line.product_in or 0, num_format=NUM_FORMAT
            )
            self._write_data_cell(
                ws, row, 4, line.product_out or 0, num_format=NUM_FORMAT
            )
            self._write_data_cell(ws, row, 5, balance, num_format=NUM_FORMAT)
